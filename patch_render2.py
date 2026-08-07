"""Патч для docx_generator.py: убирает двойную нумерацию в render_block."""
import re

with open("docx_generator.py", "r", encoding="utf-8") as f:
    content = f.read()

# 1. В render_block — убираем дублирование номера
old = '''        # Заголовок темы
        add_paragraph(doc, f"{theme_num}. {display_name}", bold=True, size=12)

        # Автор (опционально)
        if author:
            add_paragraph(doc, f"Автор: {author}")

        # Комментарий к теме (опционально)
        if theme_comment:
            add_paragraph(doc, f"Комментарий к теме: {theme_comment}")

        add_paragraph(doc, "")'''

new = '''        # Заголовок темы
        # Если display_name уже начинается с номера ("1. Смех."),
        # не добавляем ещё один номер
        if re.match(r"^\\d+\\.\\s+", display_name):
            add_paragraph(doc, display_name, bold=True, size=12)
        else:
            add_paragraph(doc, f"{theme_num}. {display_name}", bold=True, size=12)

        # Автор (опционально)
        if author:
            add_paragraph(doc, f"Автор: {author}")

        # Комментарий к теме (опционально)
        if theme_comment:
            add_paragraph(doc, f"Комментарий к теме: {theme_comment}")

        add_paragraph(doc, "")'''

if old not in content:
    print("ОШИБКА: блок заголовка темы не найден")
    raise SystemExit(1)

content = content.replace(old, new)

# 2. В списке тем в начале — тоже убираем дублирование
old2 = '''    # Список тем в начале
    add_paragraph(doc, "Список тем:", bold=True, size=12)
    for i, theme_key in enumerate(block.keys()):
        theme_num = theme_start_num + i
        display_name = get_display_name(theme_key, theme_num)
        add_paragraph(doc, f"{theme_num}. {display_name}")
    add_paragraph(doc, "")'''

new2 = '''    # Список тем в начале
    add_paragraph(doc, "Список тем:", bold=True, size=12)
    for i, theme_key in enumerate(block.keys()):
        theme_num = theme_start_num + i
        display_name = get_display_name(theme_key, theme_num)
        # Если display_name уже начинается с номера, не добавляем ещё один
        if re.match(r"^\\d+\\.\\s+", display_name):
            add_paragraph(doc, display_name)
        else:
            add_paragraph(doc, f"{theme_num}. {display_name}")
    add_paragraph(doc, "")'''

if old2 not in content:
    print("ОШИБКА: блок списка тем не найден")
    raise SystemExit(1)

content = content.replace(old2, new2)

# 3. Добавляем импорт re
old3 = '''import sys
from collections import OrderedDict

from docx import Document'''

new3 = '''import re
import sys
from collections import OrderedDict

from docx import Document'''

if old3 not in content:
    print("ОШИБКА: импорты не найдены")
    raise SystemExit(1)

content = content.replace(old3, new3)

with open("docx_generator.py", "w", encoding="utf-8") as f:
    f.write(content)

print("OK: двойная нумерация в render_block исправлена")
