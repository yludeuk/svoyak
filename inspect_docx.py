import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
d = Document('output/АС-14-Хёд-1-1..10.docx')
print('=== PARAGRAPHS ===')
for i, p in enumerate(d.paragraphs):
    style = p.style.name if p.style else 'None'
    print(f'[{i}] [{style}] {repr(p.text[:120])}')
print()
print('=== TABLES ===')
for ti, t in enumerate(d.tables):
    print(f'Table {ti}: {len(t.rows)} rows x {len(t.columns)} cols')
    for ri, row in enumerate(t.rows[:3]):
        for ci, cell in enumerate(row.cells):
            print(f'  [{ri}][{ci}] {repr(cell.text[:80])}')
