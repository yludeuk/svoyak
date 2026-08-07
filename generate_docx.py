"""Генерация 7 docx файлов из all_questions.json.

Разбивает все вопросы на блоки: 6 блоков по 10 тем + 1 блок на 12 тем = 7 файлов.
В начале каждого файла выводится список тем.
"""
import json
import sys
import os
from collections import OrderedDict
from doc_stuff import fix_settings_theme_lang, set_docdefaults_language_ru, set_normal_style_language_ru

from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")

INPUT_JSON = "all_questions.json"
OUTPUT_DIR = "output"
PACKAGE_NAME = "Проба пера"


def load_questions() -> list:
    with open(INPUT_JSON, "r", encoding="utf-8") as f:
        return json.load(f)


def group_by_theme(questions: list) -> "OrderedDict[str, list]":
    """Группирует вопросы по темам, сохраняя порядок.

    Использует (source_file, theme) как ключ, чтобы темы с одинаковым
    названием из разных файлов не объединялись.

    NULL-темы разбиваются на подтемы по 5 вопросов (так как в исходных
    файлах NULL — это секретные темы, каждая по 5 вопросов).
    """
    themes = OrderedDict()
    null_counter = 0
    null_buffer = []

    def flush_null(source_file: str):
        nonlocal null_counter, null_buffer
        if null_buffer:
            for i in range(0, len(null_buffer), 5):
                chunk = null_buffer[i : i + 5]
                null_counter += 1
                key = f"{source_file}::__secret_{null_counter}"
                themes[key] = chunk
            null_buffer = []

    current_source = None
    for q in questions:
        sf = q.get("_source_file", "?")
        if sf != current_source:
            if null_buffer:
                flush_null(current_source or "?")
            current_source = sf

        theme = q.get("theme")
        if theme is None or theme == "":
            null_buffer.append(q)
        else:
            if null_buffer:
                flush_null(current_source)
                null_buffer = []
            key = f"{sf}::{theme}"
            themes.setdefault(key, []).append(q)

    flush_null(current_source or "?")
    return themes


def split_into_blocks(themes: "OrderedDict[str, list]") -> list:
    """Разбивает на 6 блоков по 10 тем + 1 блок на оставшиеся 12 = 7 файлов."""
    all_items = list(themes.items())
    blocks = []
    for i in range(0, 60, 10):
        blocks.append(OrderedDict(all_items[i : i + 10]))
    blocks.append(OrderedDict(all_items[60:]))
    return blocks


def add_paragraph(doc: Document, text: str, bold: bool = False, size: int = 11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run.bold = bold


def get_display_name(theme_key: str, theme_num: int) -> str:
    """Возвращает отображаемое имя темы."""
    if "::" in theme_key:
        _, theme_name = theme_key.split("::", 1)
    else:
        theme_name = theme_key
    if theme_name.startswith("__secret_"):
        return f"Тема-{theme_num} (секрет)"
    return theme_name


def render_block(doc: Document, block: "OrderedDict[str, list]", block_num: int,
                 total_blocks: int, theme_start_num: int):
    """Рендерит один блок тем в документ."""
    # Заголовок документа
    add_paragraph(doc, f"Пакет «{PACKAGE_NAME}». Темы {theme_start_num}–"
                       f"{theme_start_num + len(block) - 1}.",
                  bold=True, size=14)
    add_paragraph(doc, "")

    # Список тем в начале
    add_paragraph(doc, "Список тем:", bold=True, size=12)
    for i, theme_key in enumerate(block.keys()):
        theme_num = theme_start_num + i
        display_name = get_display_name(theme_key, theme_num)
        add_paragraph(doc, f"{theme_num}. {display_name}")
    add_paragraph(doc, "")

    # Темы с вопросами
    for i, (theme_key, questions) in enumerate(block.items()):
        theme_num = theme_start_num + i
        display_name = get_display_name(theme_key, theme_num)

        add_paragraph(doc, f"{theme_num}. {display_name}", bold=True, size=12)
        add_paragraph(doc, "")

        for q in questions:
            price = q.get("price")
            question_text = q.get("question", "").strip()
            answer = q.get("answer", "").strip()
            accepted = q.get("accepted_answers")
            comment = q.get("comment")
            sources = q.get("sources")

            if price is not None:
                add_paragraph(doc, f"{price}. {question_text}")
            else:
                add_paragraph(doc, question_text)

            if answer:
                add_paragraph(doc, f"Ответ: {answer}")

            if accepted and len(accepted) > 0:
                add_paragraph(doc, f"Зачёт: {', '.join(accepted)}")

            if comment:
                add_paragraph(doc, f"Комментарий: {comment}")

            if sources and len(sources) > 0:
                if len(sources) == 1:
                    add_paragraph(doc, f"Источник: {sources[0]}")
                else:
                    add_paragraph(doc, f"Источники: {', '.join(sources)}")

            add_paragraph(doc, "")

        add_paragraph(doc, "")


def main():
    if not os.path.exists(INPUT_JSON):
        print(f"[!] Файл {INPUT_JSON} не найден")
        sys.exit(1)

    os.makedirs(OUTPUT_DIR, exist_ok=True)

    print(f"[*] Загружаю {INPUT_JSON}...")
    questions = load_questions()
    print(f"[*] Всего вопросов: {len(questions)}")

    print("[*] Группирую по темам...")
    themes = group_by_theme(questions)
    print(f"[*] Всего тем: {len(themes)}")

    blocks = split_into_blocks(themes)
    print(f"[*] Получено блоков: {len(blocks)}")
    for i, b in enumerate(blocks, 1):
        print(f"    Блок {i}: {len(b)} тем")

    theme_counter = 1
    for i, block in enumerate(blocks, 1):
        first = theme_counter
        last = theme_counter + len(block) - 1
        fname = f"{PACKAGE_NAME}-{i}-{first}..{last}.docx"
        fpath = os.path.join(OUTPUT_DIR, fname)

        doc = Document()
        doc = Document()
        fix_settings_theme_lang(doc)
        set_docdefaults_language_ru(doc)
        set_normal_style_language_ru(doc)
        section = doc.sections[0]
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        render_block(doc, block, i, len(blocks), first)
        doc.save(fpath)
        print(f"[*] Сохранён {fpath}")
        theme_counter = last + 1

    print(f"\n[*] Готово! Создано {len(blocks)} файлов в {OUTPUT_DIR}/")


if __name__ == "__main__":
    main()
