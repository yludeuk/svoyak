"""Патч для docx_generator.py: убирает двойную нумерацию в списке тем."""
import re

with open("docx_generator.py", "r", encoding="utf-8") as f:
    content = f.read()

# Исправляем список тем в начале — не добавляем theme_num, если display_name уже с номером
old = '''    # Список тем в начале
    add_paragraph(doc, "Список тем:", bold=True, size=12)
    for i, theme_key in enumerate(block.keys()):
        theme_num = theme_start_num + i
        display_name = get_display_name(theme_key, theme_num)
        add_paragraph(doc, f"{theme_num}. {display_name}")
    add_paragraph(doc, "")'''

new = '''    # Список тем в начале
    add_paragraph(doc, "Список тем:", bold=True, size=12)
    for i, theme_key in enumerate(block.keys()):
        theme_num = theme_start_num + i
        display_name = get_display_name(theme_key, theme_num)
        # Если display_name уже начинается с номера, не добавляем ещё один
        if re.match(r"^\\d+\\.\\s+", display_name):
            add_paragraph(doc, display_name)
        else:
            add_paragraph(doc, f"{theme_num}. {display_name}")
    add_paragraph(doc, "")'''

if old not in content:
    print("ОШИБКА: блок списка тем не найден")
    raise SystemExit(1)

content = content.replace(old, new)

# Добавляем импорт re если его нет
if "import re" not in content:
    old2 = '''import sys
from collections import OrderedDict

from docx import Document'''
    new2 = '''import re
import sys
from collections import OrderedDict

from docx import Document'''
    if old2 not in content:
        print("ОШИБКА: импорты не найдены")
        raise SystemExit(1)
    content = content.replace(old2, new2)

with open("docx_generator.py", "w", encoding="utf-8") as f:
    f.write(content)

print("OK: двойная нумерация в списке тем исправлена")
